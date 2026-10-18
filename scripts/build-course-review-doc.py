"""Create a US-letter review worksheet from the current course review notes."""
from pathlib import Path
import json
import re
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

root=Path(__file__).resolve().parent.parent
source=(root/'courses/REMAINING_LAUNCH_REVIEW.md').read_text()
catalog=json.loads((root/'courses/approved-launch-catalog.json').read_text())
titles={c['code']:c['title'] for a in catalog['academies'] for c in a['courses']}
intro,tail=source.split('## Items to review before publishing',1)
priorities,per_course=tail.split('## Per-course notes',1)
courses=[]
for chunk in re.split(r'^### ',per_course,flags=re.M)[1:]:
    code,body=chunk.split('\n',1)
    courses.append((code.strip(),[line[2:].strip() for line in body.splitlines() if line.startswith('- ')]))
assert len(courses)==21

doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.7)
section.header_distance=section.footer_distance=Inches(.3)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.05
for name,size in [('Title',23),('Heading 1',16),('Heading 2',13)]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('172747')
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.space_before=Pt(6)
header=section.header.paragraphs[0]
header.text='colTRAIN  |  Course review worksheet'
header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('595959')
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('colTRAIN • Draft review  |  Page ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)


def rich(paragraph,text):
    for i,part in enumerate(re.split(r'(\*\*.*?\*\*)',text)):
        run=paragraph.add_run(part[2:-2] if part.startswith('**') else part)
        run.bold=part.startswith('**')


def ruled_line(parent):
    p=parent.add_paragraph();p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=Pt(21)
    p.add_run(' ')
    borders=OxmlElement('w:pBdr')
    for edge in ['bottom','between']:
        border=OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','3'),('color','B7B7B7'),('space','1')]:border.set(qn('w:'+key),value)
        borders.append(border)
    p._p.get_or_add_pPr().append(borders)


doc.add_heading('Course review notes',0)
doc.add_paragraph('Remaining 21 launch courses • Package version 1.0.0')
doc.add_paragraph('Reviewer: __________________________     Review date: __________________')
for paragraph in intro.split('\n\n'):
    if paragraph.strip() and not paragraph.startswith('# '):rich(doc.add_paragraph(),paragraph.strip())
doc.add_heading('Items to review before publishing',1)
for line in priorities.splitlines():
    if line.startswith('- '):rich(doc.add_paragraph(style='List Bullet'),line[2:])
doc.add_paragraph('Use the following course sheets to record wording changes, questions, references to verify, and media decisions. The notes are copied from the course review; course titles and writing space have been added for convenience.')

for start in range(0,len(courses),2):
    doc.add_page_break()
    table=doc.add_table(rows=min(2,len(courses)-start),cols=1)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False;table.columns[0].width=Inches(7.1)
    for row,(code,notes) in zip(table.rows,courses[start:start+2]):
        row.height=Inches(4.65);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
        cell=row.cells[0];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        p=cell.paragraphs[0];p.style=doc.styles['Heading 2'];p.add_run(code+' — '+titles[code])
        for note in notes:rich(cell.add_paragraph(style='List Bullet'),note)
        p=cell.add_paragraph('Your notes / proposed changes:');p.runs[0].bold=True
        for _ in range(4):ruled_line(cell)
        p=cell.add_paragraph('☐ Needs changes     ☐ Ready for pilot review     ☐ Hold for references/media')
        p.paragraph_format.space_before=Pt(8)
        for run in p.runs:run.font.size=Pt(9)
        p=cell.add_paragraph('Follow-up / owner: __________________________     Date: ______________')
        for run in p.runs:run.font.size=Pt(9)

doc.core_properties.title='colTRAIN — Course Review Notes'
doc.core_properties.subject='Printable review worksheet for the remaining 21 launch courses'
doc.core_properties.author='colTRAIN'
doc.core_properties.keywords='course review, draft, worksheet, launch catalog'
output=root/'design/colTRAIN_Course_Review_Notes.docx';doc.save(output)
check=Document(output)
all_text='\n'.join([p.text for p in check.paragraphs]+[p.text for t in check.tables for row in t.rows for cell in row.cells for p in cell.paragraphs])
for code,notes in courses:
    assert code+' — '+titles[code] in all_text
    for note in notes:assert note.replace('**','') in all_text
print(f'Created {output.name}: all 21 courses and every review note, US Letter, two courses per worksheet page.')
